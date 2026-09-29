import os
import re
import time
import logging
from datetime import datetime, date
from io import BytesIO
from functools import wraps

import pandas as pd
from sqlalchemy import create_engine, text

from config import Config

logger = logging.getLogger(__name__)

engine = create_engine(Config.DATABASE_URL)


# ==================== ДЕКОРАТОР ЛОГИРОВАНИЯ МЕДЛЕННЫХ ЗАПРОСОВ ====================
def log_slow_query(func):
    """
    Декоратор для функций, которые выполняют SQL-запросы.
    Логирует время выполнения; если превышает SLOW_QUERY_THRESHOLD — warning.
    """
    @wraps(func)
    def wrapper(*args, **kwargs):
        start = time.time()
        result = func(*args, **kwargs)
        elapsed = time.time() - start
        name = func.__name__
        if elapsed >= Config.SLOW_QUERY_THRESHOLD:
            logger.warning(f"SLOW QUERY: {name} выполнился за {elapsed:.2f} сек")
        else:
            logger.debug(f"{name} выполнился за {elapsed:.2f} сек")
        return result
    return wrapper


# ==================== ОБЩИЕ ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ====================
def normalize_doc_ref(text):
    """
    Нормализует строку 'Документ расчётов' для сопоставления КЗ ↔ заявок.
    Убирает лишние пробелы, время, приводит к нижнему регистру.
    Сохраняет уникальность: номер + дата остаются.
    """
    if pd.isna(text) or str(text).strip() == '':
        return None
    s = str(text)
    # Неразрывные и «тонкие» пробелы → обычный
    s = s.replace('\xa0', ' ').replace('\u202f', ' ').replace('\u2009', ' ')
    # Убираем время (HH:MM:SS или HH:MM)
    s = re.sub(r'\s+\d{1,2}:\d{2}(?::\d{2})?', '', s)
    # Приводим к нижнему регистру и сжимаем пробелы
    s = s.lower()
    s = re.sub(r'\s+', ' ', s).strip()
    return s if s else None


def split_doc_refs(text):
    """
    Разбирает строку с несколькими документами (через ';') в список нормализованных ключей.
    """
    if pd.isna(text) or str(text).strip() == '':
        return []
    parts = [p.strip() for p in str(text).split(';') if p.strip()]
    return [k for k in (normalize_doc_ref(p) for p in parts) if k]

def normalize_date(date_val, dayfirst=True):
    """Преобразует различные представления дат в объект datetime.date."""
    if pd.isna(date_val):
        return None
    if isinstance(date_val, pd.Timestamp):
        return date_val.date()
    try:
        return pd.to_datetime(date_val, dayfirst=dayfirst).date()
    except Exception:
        return None


def clean_money(value):
    """Очищает денежное значение: убирает пробелы, заменяет ',' на '.'."""
    if pd.isna(value):
        return 0.0
    s = str(value).replace(' ', '').replace(',', '.')
    try:
        return float(s)
    except Exception:
        return 0.0


def parse_date_dd_mm_yyyy(s):
    """Парсит дату формата DD.MM.YYYY или DD.MM.YY."""
    if pd.isna(s) or s == '':
        return None
    for fmt in ('%d.%m.%Y', '%d.%m.%y'):
        try:
            return datetime.strptime(str(s).strip(), fmt).date()
        except Exception:
            pass
    return None


# ==================== ЭКСПОРТ В EXCEL С ФОРМАТИРОВАНИЕМ ====================
def _sanitize_cell(value):
    """
    Приводит значение ячейки к типу, который openpyxl может записать в Excel.
    - list/tuple/set → строка через '; '
    - dict → JSON-строка
    - None/NaN → пустая строка
    - datetime/date → как есть (openpyxl сам обработает)
    - остальное → как есть
    """
    import json
    from datetime import date, datetime

    # NaN / None
    try:
        if pd.isna(value):
            return ''
    except (TypeError, ValueError):
        pass

    if value is None:
        return ''

    # Даты и числа — оставляем как есть
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, str):
        return value

    # Списки, кортежи, множества
    if isinstance(value, (list, tuple, set)):
        if not value:
            return ''
        try:
            return '; '.join(str(v) for v in value)
        except Exception:
            return str(value)

    # Словари
    if isinstance(value, dict):
        if not value:
            return ''
        try:
            return json.dumps(value, ensure_ascii=False)
        except Exception:
            return str(value)

    # Всё остальное — строкой
    return str(value)


def export_df_with_formatting(df, sheet_name='Отчёт'):
    """
    Формирует Excel-файл из DataFrame с:
      - оформлением в виде таблицы (аналог Ctrl+T);
      - применением числовых форматов к столбцам;
      - автошириной столбцов;
      - нормализацией сложных значений (list/dict) в строки.
    Возвращает BytesIO.
    """
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    # Санитизация всех ячеек ДО записи
    df_clean = df.copy()
    for col in df_clean.columns:
        df_clean[col] = df_clean[col].apply(_sanitize_cell)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel ограничивает 31 символами

    headers = [str(c) for c in df_clean.columns]
    ws.append(headers)

    for row in df_clean.itertuples(index=False, name=None):
        ws.append(list(row))

    n_rows = len(df_clean) + 1
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)

    table_name = 'ReportTable'
    table = Table(displayName=table_name, ref=f"A1:{last_col_letter}{n_rows}")
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    # Форматы столбцов
    MONEY_COLS = {
        'сумма_заявки', 'сумма_пп', 'сумма_согласованная', 'сумма_заявленная',
        'факт_исполнения', 'amount', 'spent', 'kz', 'dz', 'unpaid_sum',
        'сумма_остаток', 'остаток', 'kz_sum', 'apps_sum', 'diff', 'total',
        'apps_count',
    }
    MONTH_YEAR_COLS = {'период_услуги'}
    DATE_COLS = {'дата', 'дата_реестра', 'дата_отправки', 'дата_принятия_банком',
                 'дата_оплаты_фактич', 'start_date', 'end_date',
                 'дата_документа', 'период_погашения'}

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        col_lower = col_name.lower()

        if col_lower in MONEY_COLS:
            for r in range(2, n_rows + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.number_format = '#,##0.00'
                if cell.value is not None and cell.value != '':
                    try:
                        cell.value = float(cell.value)
                    except (TypeError, ValueError):
                        pass

        elif col_lower in MONTH_YEAR_COLS:
            for r in range(2, n_rows + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.number_format = 'MMMM YYYY'

        elif col_lower in DATE_COLS:
            for r in range(2, n_rows + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.number_format = 'DD.MM.YYYY'

        # Автоширина
        max_len = len(str(col_name))
        for r in range(2, n_rows + 1):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ==================== СОЗДАНИЕ ТАБЛИЦ ====================
def create_all_tables(engine):
    from modules.contracts.models import create_contracts_tables

    with engine.connect() as conn:
        # ---------- applications ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS applications (
            номер_заявки VARCHAR(50),
            дата DATE,
            цфо TEXT,
            проект TEXT,
            статья_оборотов TEXT,
            вид_затрат TEXT,
            стг_подсказка_исд TEXT,
            код TEXT,
            контрагент TEXT,
            признак_оплаты TEXT,
            договор_контрагента TEXT,
            договор_код TEXT,
            документ_расчетов_с_контрагентом TEXT,
            назначение_платежа TEXT,
            номер_рп TEXT,
            оплачена VARCHAR(3),
            платежное_поручение_исходящее TEXT,
            статус_согласования TEXT,
            дата_последнего_статуса TIMESTAMP,
            замечания_дф TEXT,
            состояние_заявки TEXT,
            ответственный TEXT,
            сумма_заявки NUMERIC(18,2),
            сумма_пп NUMERIC(18,2),
            период_услуги DATE,
            дата_оплаты_фактич DATE,
            is_bezreestroviy BOOLEAN DEFAULT FALSE,
            PRIMARY KEY (номер_заявки, дата)
        );
        """))

        # Миграции для существующих баз (идемпотентно)
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS "
            "документ_расчетов_с_контрагентом TEXT;"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS "
            "дата_оплаты_фактич DATE;"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS "
            "is_bezreestroviy BOOLEAN DEFAULT FALSE;"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS "
            "назначение_платежа TEXT;"
        ))

        # ---------- registers ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS registers (
            id SERIAL PRIMARY KEY,
            номер_реестра TEXT UNIQUE,
            номер_реестра_исходный TEXT,
            дата_реестра DATE,
            дата_отправки DATE,
            дата_поступления_ксд DATE,
            дата_поступления_кд DATE,
            дата_поступления_кбс DATE,
            дата_принятия_банком DATE,
            реестр_зарплатных TEXT,
            вид_реестра TEXT,
            тип_затрат TEXT,
            тип_платежей TEXT,
            инициатор TEXT,
            наименование_инициатора TEXT,
            наименование_плательщика TEXT,
            получатель TEXT,
            подал_ксд TEXT,
            факт_исполнения NUMERIC(18,2),
            сумма_заявленная NUMERIC(18,2),
            сумма_согласованная NUMERIC(18,2),
            статус TEXT,
            дополнительный_статус TEXT,
            id_заявки TEXT,
            исд TEXT
        );
        """))

        # ---------- debts ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS debts (
            id SERIAL PRIMARY KEY,
            вид_задолженности TEXT,
            счет TEXT,
            архивный_номер TEXT,
            контрагент TEXT,
            инн TEXT,
            состояние_контрагента TEXT,
            договор_код TEXT,
            договор_вид TEXT,
            договор_характер TEXT,
            договор_наименование TEXT,
            проект TEXT,
            цфо TEXT,
            номер_документа TEXT,
            дата_документа DATE,
            документ_расчетов TEXT,
            период_погашения DATE,
            дней_просрочки INTEGER,
            сумма_остаток NUMERIC(18,2),
            imported_at TIMESTAMP DEFAULT NOW()
        );
        """))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_debts_договор_код ON debts(договор_код);"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_debts_контрагент ON debts(контрагент);"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_debts_документ_расчетов ON debts(документ_расчетов);"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS договор_код TEXT;"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_applications_договор_код "
            "ON applications(договор_код);"
        ))
        # ---------- Токены объектов (для этапа 5) ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS object_tokens_raw (
            id SERIAL PRIMARY KEY,
            token TEXT NOT NULL,
            source_field TEXT NOT NULL,
            source_id TEXT,
            occurrences INTEGER DEFAULT 1,
            examples JSONB DEFAULT '[]'::jsonb,
            created_at TIMESTAMP DEFAULT NOW()
        );
        """))
        # ---------- Сохранённые отчёты конструктора ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS saved_reports (
            id SERIAL PRIMARY KEY,
            name TEXT NOT NULL,
            description TEXT,
            report_type TEXT NOT NULL DEFAULT 'table',
            config JSONB NOT NULL,
            owner_id INTEGER REFERENCES users(id),
            is_public BOOLEAN DEFAULT FALSE,
            is_system BOOLEAN DEFAULT FALSE,
            created_at TIMESTAMP DEFAULT NOW(),
            updated_at TIMESTAMP DEFAULT NOW()
        );
        """))
        conn.execute(text("""
        CREATE INDEX IF NOT EXISTS idx_saved_reports_owner ON saved_reports(owner_id);
        """))
        # ---------- Персональные дашборды ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS user_dashboards (
            id SERIAL PRIMARY KEY,
            user_id INTEGER REFERENCES users(id) ON DELETE CASCADE UNIQUE,
            widgets JSONB DEFAULT '[]'::jsonb,
            is_default BOOLEAN DEFAULT TRUE,
            updated_at TIMESTAMP DEFAULT NOW()
        );
        """))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_user_dashboards_user ON user_dashboards(user_id);"
        ))
        # ---------- Персональные настройки sidebar ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS user_sidebar (
            user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
            pinned_sections JSONB DEFAULT '[]'::jsonb,
            hidden_sections JSONB DEFAULT '[]'::jsonb,
            updated_at TIMESTAMP DEFAULT NOW()
        );
        """))
        # ---------- Справочник объектов ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS objects (
            id SERIAL PRIMARY KEY,
            code TEXT UNIQUE,                -- 3-значный код для заявок
            name TEXT NOT NULL,
            short_name TEXT,
            type TEXT,                       -- res / site / office
            status TEXT,                     -- active / delayed / stopped
            project TEXT,                    -- проект
            field TEXT,                      -- месторождение (справочно)
            lat NUMERIC(10,7),
            lon NUMERIC(10,7),
            ответственный TEXT,
            info TEXT,
            created_at TIMESTAMP DEFAULT NOW()
        );
        """))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_objects_code ON objects(code);"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_objects_project ON objects(project);"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_objects_type ON objects(type);"
        ))
        conn.execute(text("""
        CREATE INDEX IF NOT EXISTS idx_obj_tokens_token ON object_tokens_raw(token);
        """))
        conn.execute(text("""
        CREATE INDEX IF NOT EXISTS idx_obj_tokens_source ON object_tokens_raw(source_field);
        """))

        # ---------- Справочник согласующих по проектам ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS approval_projects (
            id SERIAL PRIMARY KEY,
            isd TEXT,
            project TEXT NOT NULL,
            platform TEXT,
            registry_type TEXT,
            created_at TIMESTAMP DEFAULT NOW(),
            updated_at TIMESTAMP DEFAULT NOW(),
            UNIQUE (isd, project)
        );
        """))
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS approvers (
            id SERIAL PRIMARY KEY,
            name TEXT NOT NULL,
            email TEXT,
            phone TEXT,
            organization TEXT,
            created_at TIMESTAMP DEFAULT NOW(),
            updated_at TIMESTAMP DEFAULT NOW()
        );
        """))
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS project_stage_approvers (
            id SERIAL PRIMARY KEY,
            project_id INTEGER REFERENCES approval_projects(id) ON DELETE CASCADE,
            stage INTEGER NOT NULL CHECK (stage IN (1, 2)),
            approver_id INTEGER REFERENCES approvers(id) ON DELETE CASCADE,
            order_num INTEGER DEFAULT 0,
            UNIQUE (project_id, stage, approver_id)
        );
        """))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_approval_projects_isd "
            "ON approval_projects(isd);"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_approval_projects_project "
            "ON approval_projects(project);"
        ))
        conn.execute(text(
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_approvers_name_email "
            "ON approvers (name, COALESCE(email, ''));"
        ))
        # Флаг прав редактирования справочника
        conn.execute(text(
            "ALTER TABLE users ADD COLUMN IF NOT EXISTS can_edit_approvers "
            "BOOLEAN DEFAULT FALSE;"
        ))

        # ---------- Токены объектов + матчинг ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS object_tokens (
            object_id INTEGER REFERENCES objects(id) ON DELETE CASCADE,
            token TEXT NOT NULL,
            weight NUMERIC(6,4),
            PRIMARY KEY (object_id, token)
        );
        """))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_object_tokens_token ON object_tokens(token);"
        ))
        # Столбцы для будущей записи матчей (пока не заполняются)
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS object_id INTEGER;"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS object_match_confidence TEXT;"
        ))
        conn.execute(text(
            "ALTER TABLE applications ADD COLUMN IF NOT EXISTS object_match_candidates JSONB;"
        ))
        # ---------- users ----------
        conn.execute(text("""
        CREATE TABLE IF NOT EXISTS users (
            id SERIAL PRIMARY KEY,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'viewer',
            цфо TEXT,
            can_import BOOLEAN DEFAULT FALSE,
            must_change_password BOOLEAN DEFAULT TRUE,
            is_active BOOLEAN DEFAULT TRUE,
            created_at TIMESTAMP DEFAULT NOW(),
            last_login TIMESTAMP
        );
        """))
        conn.execute(text(
            "ALTER TABLE users ADD COLUMN IF NOT EXISTS start_page TEXT DEFAULT 'home';"
        ))
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS idx_users_email ON users(LOWER(email));"
        ))
        # ---------- contracts + application_contracts ----------
        create_contracts_tables(conn)

        conn.commit()


# ==================== СОЗДАНИЕ ПРЕДСТАВЛЕНИЙ ====================
def create_all_views(engine):
    sql = """
    DROP VIEW IF EXISTS v_pq_quick_report CASCADE;
    DROP VIEW IF EXISTS v_applications_full CASCADE;
    DROP VIEW IF EXISTS v_registers CASCADE;

    CREATE VIEW v_registers AS
    SELECT *, COALESCE(дополнительный_статус, статус) AS вычисленный_статус
    FROM registers;

    CREATE VIEW v_applications_full AS
    SELECT 
        a.*,
        r.id AS id_реестра,
        r.дата_реестра,
        r.статус AS статус_реестра,
        r.дополнительный_статус AS доп_статус_реестра,
        COALESCE(r.дополнительный_статус, r.статус) AS итоговый_статус_реестра,
        r.сумма_заявленная,
        r.сумма_согласованная,
        r.факт_исполнения,
        r.дата_принятия_банком,
        r.получатель AS получатель_по_реестру
    FROM applications a
    LEFT JOIN v_registers r ON a.номер_рп = r.номер_реестра;

    -- v_pq_quick_report оставляем как отдельное представление для совместимости,
    -- но быстрый отчёт в приложении строится прямым SQL-запросом из v_applications_full.
    CREATE VIEW v_pq_quick_report AS
    SELECT 
        a.номер_заявки,
        a.код AS код_проекта,
        a.контрагент,
        a.сумма_заявки,
        a.номер_рп AS номер_реестра,
        CASE 
            WHEN a.is_bezreestroviy AND a.итоговый_статус_реестра IS NULL THEN 'Безреестровый'
            ELSE a.итоговый_статус_реестра
        END AS итоговый_статус_реестра,
        a.проект,
        a.статус_согласования,
        a.период_услуги
    FROM v_applications_full a
    WHERE a.оплачена = 'Нет'
      AND a.статус_согласования IN ('Согласовано ДФ', 'Согласовано руководителем ЦФО')
      AND (a.проект = 'СТГ-00' OR a.номер_рп IS NOT NULL);
    """
    with engine.connect() as conn:
        conn.execute(text(sql))
        conn.commit()

def current_cfo():
    """
    Возвращает ЦФО текущего пользователя для фильтрации данных.
    None означает «фильтр не применять» — для админов и пользователей без ЦФО.
    """
    try:
        from flask_login import current_user
        if not current_user.is_authenticated:
            return None
        if getattr(current_user, 'is_admin', False):
            return None
        dep = getattr(current_user, 'department', None)
        return dep if dep else None
    except Exception:
        return None

def read_sql(query, *args, params=None, **kwargs):
    """
    Безопасная обёртка над pd.read_sql.
    - Игнорирует переданный engine — использует глобальный.
    - Работает с psycopg2-стилем параметров %(name)s (не :name).
    - Корректно обрабатывает пустые params.
    - Возвращает DataFrame с колонками, даже если строк нет.
    """
    with engine.connect() as conn:
        if params:
            result = conn.exec_driver_sql(query, params)
        else:
            result = conn.exec_driver_sql(query)

        rows = result.fetchall()
        columns = list(result.keys()) if result.keys() else []

        if not rows:
            return pd.DataFrame(columns=columns)
        return pd.DataFrame(rows, columns=columns)